# src/results/service.py
import os
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.section import WD_ORIENTATION
from datetime import datetime
from typing import List, Dict, Any


class WordExporter:
    """Экспортер результатов в Word документ"""

    @staticmethod
    def set_cell_border(cell, border_size=1):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        for edge in ['top', 'left', 'bottom', 'right']:
            edge_element = OxmlElement(f'w:{edge}')
            edge_element.set(qn('w:val'), 'single')
            edge_element.set(qn('w:sz'), str(border_size))
            edge_element.set(qn('w:space'), '0')
            edge_element.set(qn('w:color'), '000000')
            tcPr.append(edge_element)

    @staticmethod
    def set_page_landscape(doc):
        section = doc.sections[0]
        section.orientation = WD_ORIENTATION.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    @staticmethod
    def add_header(doc, competition_name: str):
        """Добавляет верхний колонтитул с названием соревнования на каждой странице"""
        section = doc.sections[0]
        header = section.header
        header_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        header_para.text = f"Соревнования \"{competition_name}\""
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_para.style = doc.styles['Header']
        for run in header_para.runs:
            run.font.size = Pt(10)
            run.font.bold = True

    @staticmethod
    def _get_gender_string(gender):
        if gender is None:
            return ''
        if hasattr(gender, 'value'):
            return 'М' if gender.value == 1 else 'Ж'   # предположим MALE=1, FEMALE=2
        return str(gender)[0] if str(gender) else ''

    @staticmethod
    def _get_date_string(date_value):
        if date_value is None:
            return ''
        if hasattr(date_value, 'strftime'):
            return date_value.strftime('%d.%m.%Y')
        return str(date_value)

    @staticmethod
    def export_competition_results_detailed(
            results: List[Dict[str, Any]],
            competition_name: str,
            competition_date: str,
            competition_location: str = None,
            organizer: str = None,
            total_participants: int = None,
            male_count: int = None,
            female_count: int = None
    ) -> str:
        doc = Document()
        WordExporter.set_page_landscape(doc)
        WordExporter.add_header(doc, competition_name)

        style = doc.styles['Normal']
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)

        # Заголовок документа
        title = doc.add_heading('ПРОТОКОЛ', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

        # Информация о соревнованиях
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f'соревнований по армрестлингу\n"{competition_name}"')
        run.bold = True
        run.font.size = Pt(14)

        doc.add_paragraph()

        # Информационная таблица
        info_table = doc.add_table(rows=4, cols=2)
        info_table.style = 'Table Grid'
        info_table.autofit = False
        info_table.allow_autofit = False
        info_table.columns[0].width = Cm(5)
        info_table.columns[1].width = Cm(21)

        info_data = [
            ('Дата проведения:', competition_date),
            ('Место проведения:', competition_location or 'не указано'),
            ('Организатор:', organizer or 'не указан'),
        ]
        for i, (label, value) in enumerate(info_data):
            if i < len(info_table.rows):
                info_table.rows[i].cells[0].text = label
                info_table.rows[i].cells[1].text = value
                info_table.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True
                for cell in info_table.rows[i].cells:
                    WordExporter.set_cell_border(cell)

        # Добавить строку с количеством участников и разбивкой по полу
        if total_participants is not None and male_count is not None and female_count is not None:
            stats_row = info_table.add_row()
            stats_row.cells[0].text = 'Участники:'
            stats_row.cells[1].text = f'Всего: {total_participants}, мужчин: {male_count}, женщин: {female_count}'
            stats_row.cells[0].paragraphs[0].runs[0].font.bold = True
            for cell in stats_row.cells:
                WordExporter.set_cell_border(cell)

        doc.add_paragraph()

        # Далее идут возрастные категории и таблицы (как в оригинале)
        for age_group in results:
            doc.add_heading(f'{age_group["age_category"]}', level=1)
            for weight_group in age_group["weight_categories"]:
                doc.add_heading(f'Весовая категория: {weight_group["weight_category"]}', level=2)

                table = doc.add_table(rows=1, cols=10)
                table.style = 'Table Grid'
                table.autofit = False
                table.allow_autofit = False
                widths = [Cm(1.2), Cm(3.5), Cm(1.5), Cm(2.5),
                          Cm(1.5), Cm(1.3), Cm(1.5), Cm(1.3), Cm(1.3), Cm(1.2)]
                for i, width in enumerate(widths):
                    if i < len(table.columns):
                        table.columns[i].width = width

                headers = ['Место', 'ФИО', 'Разряд', 'Команда',
                           'Левая\nрука\n(место)', 'Левая\nрука\n(очки)',
                           'Правая\nрука\n(место)', 'Правая\nрука\n(очки)',
                           'Сумма\nочков', 'Вес, кг']
                header_cells = table.rows[0].cells
                for i, header in enumerate(headers):
                    if i < len(header_cells):
                        header_cells[i].text = header
                        header_cells[i].paragraphs[0].runs[0].font.bold = True
                        header_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        WordExporter.set_cell_border(header_cells[i])

                participants = sorted(weight_group["participants"], key=lambda x: x.get('place_by_two_arms', 999))
                for participant in participants:
                    row_cells = table.add_row().cells
                    athlete_weight = participant.get('athlete_weight')
                    weight_str = f"{athlete_weight}" if athlete_weight else ''
                    cell_values = [
                        str(participant.get('place_by_two_arms', '')),
                        participant.get('full_name', ''),
                        participant.get('rank', '') or '',
                        participant.get('team', '') or '',
                        str(participant.get('left_hand_place', '')),
                        str(participant.get('left_points', 0)),
                        str(participant.get('right_hand_place', '')),
                        str(participant.get('right_points', 0)),
                        str(participant.get('total_points', 0)),
                        weight_str
                    ]
                    for i, value in enumerate(cell_values):
                        if i < len(row_cells):
                            row_cells[i].text = value
                            row_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                            WordExporter.set_cell_border(row_cells[i])
                doc.add_paragraph()

        # Подписи
        doc.add_paragraph()
        doc.add_paragraph()
        signatures = doc.add_table(rows=2, cols=1)
        signatures.autofit = False
        if len(signatures.columns) >= 1:
            signatures.columns[0].width = Cm(28)
        if len(signatures.rows) >= 2:
            signatures.rows[0].cells[0].text = 'Главный судья'
            signatures.rows[1].cells[0].text = ''
            for row in signatures.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    WordExporter.set_cell_border(cell)

        os.makedirs('exports', exist_ok=True)
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        safe_name = competition_name.replace('/', '_').replace('\\', '_').replace(':', '_')
        output_path = f'exports/protocol_{safe_name}_{timestamp}.docx'
        doc.save(output_path)
        return output_path